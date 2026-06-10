import os
import csv
import io
import logging

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
MAX_TEXT_LENGTH = 8000


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".txt":
            return _read_txt(file_path)
        elif ext == ".csv":
            return _read_csv(file_path)
        elif ext == ".pdf":
            return _read_pdf(file_path)
        elif ext == ".docx":
            return _read_docx(file_path)
        elif ext == ".pptx":
            return _read_pptx(file_path)
        elif ext in (".xls", ".xlsx"):
            return _read_excel(file_path)
        elif ext in IMAGE_EXTENSIONS:
            return ""
        else:
            logger.info(f"No text extractor for {ext}")
            return ""
    except Exception as e:
        logger.warning(f"Extraction failed for {file_path}: {e}")
        return ""


def _read_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()[:MAX_TEXT_LENGTH]


def _read_csv(path: str) -> str:
    rows = []
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            if i > 100:
                rows.append("... (truncated)")
                break
            rows.append(" | ".join(row))
    text = "\n".join(rows)
    return text[:MAX_TEXT_LENGTH]


def _read_pdf(path: str) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            if i >= 20:
                break
            t = page.extract_text()
            if t:
                pages.append(t.strip())
    text = "\n\n".join(pages)
    return text[:MAX_TEXT_LENGTH]


def _read_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paras)
    return text[:MAX_TEXT_LENGTH]


def _read_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        parts.append(para.text.strip())
    text = "\n".join(parts)
    return text[:MAX_TEXT_LENGTH]


def _read_excel(path: str) -> str:
    import pandas as pd
    dfs = pd.read_excel(path, sheet_name=None)
    lines = []
    for sheet_name, df in dfs.items():
        lines.append(f"--- Sheet: {sheet_name} ---")
        lines.append(df.head(50).to_string(index=False))
    text = "\n".join(lines)
    return text[:MAX_TEXT_LENGTH]
